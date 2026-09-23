# ===============================================================================
# Copyright 2026 ross
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
# ===============================================================================
"""Tests for the chemistry field-sheet ingest
(services/chemistry_field_params.py)."""

from datetime import datetime

import pytest
from openpyxl import Workbook
from sqlalchemy import delete, select

from db.engine import session_ctx
from db.nma_legacy import NMA_Chemistry_SampleInfo, NMA_FieldParameters
from services.chemistry_field_params import (
    FIELD_SHEET_DATASET,
    FieldParamsMappingError,
    import_field_tables,
    prep_field_parameters,
    prep_sample_info,
    replay_field_sheet,
    select_tables,
    upload_field_export,
)
from services.ingest_raw_zone import list_snapshots, read_snapshot
from services.chemistry_field_sheet import SheetTable, read_local_export

WELL = "Test Well"

SAMPLE_INFO_HEADER = [
    "WellPointID",
    "SamplePointID",
    "AnalysisAgency",
    "SampleType (SD)",
    "CollectionDate",
    "CollectionMethod (F = faucet)",
    "CollectedBy (GR = grab??)",
    "Data Source",
    "Staff",
    "Sample Notes",
]

FIELD_PARAMS_HEADER = [
    "SamplePointID",
    "Time",
    "Discharge rate (gpm)",
    "pHf",
    "T (C)",
    "CF (uS/cm)",
    "DO (mg/L)",
    "ORP (mV)",
]


def _sample_info_row(**overrides):
    row = {
        "WellPointID": WELL,
        "SamplePointID": f"{WELL}A",
        "AnalysisAgency": "NMBGMR",
        "SampleType (SD)": "SD",
        "CollectionDate": "2025-06-10T12:05:00",
        "CollectionMethod (F = faucet)": "Faucet at well head",
        "CollectedBy (GR = grab??)": None,
        "Data Source": "NMBGMR",
        "Staff": "Dan Lavery, Sianin Spaur",
        "Sample Notes": "Sampled from spigot",
    }
    row.update(overrides)
    return row


def _field_params_row(**overrides):
    row = {
        "SamplePointID": f"{WELL}A",
        "Time": "2025-06-10T12:07:00",
        "Discharge rate (gpm)": 0.568,
        "pHf": 6.94,
        "T (C)": 15.7,
        "CF (uS/cm)": 783.0,
        "DO (mg/L)": 1.35,
        "ORP (mV)": 78.5,
    }
    row.update(overrides)
    return row


def _table(title: str, header: list[str], rows: list[dict]) -> SheetTable:
    records = []
    for offset, row in enumerate(rows):
        record = {column: row.get(column) for column in header}
        record[SheetTable.ROW_NUMBER_KEY] = offset + 2
        records.append(record)
    return SheetTable(title=title, header=header, rows=records)


def _tables(sample_rows=None, param_rows=None) -> list[SheetTable]:
    return [
        _table(
            "ChemistrySampleInfo",
            SAMPLE_INFO_HEADER,
            [_sample_info_row()] if sample_rows is None else sample_rows,
        ),
        _table(
            "FieldParameters",
            FIELD_PARAMS_HEADER,
            [_field_params_row()] if param_rows is None else param_rows,
        ),
    ]


@pytest.fixture()
def _cleanup_field_chemistry():
    """Remove samples (and cascaded field parameters) created during a test."""
    yield
    with session_ctx() as session:
        session.execute(
            delete(NMA_Chemistry_SampleInfo).where(
                NMA_Chemistry_SampleInfo.nma_sample_point_id.like(f"{WELL}%")
            )
        )
        session.commit()


def _samples():
    with session_ctx() as session:
        return session.scalars(
            select(NMA_Chemistry_SampleInfo)
            .where(NMA_Chemistry_SampleInfo.nma_sample_point_id.like(f"{WELL}%"))
            .order_by(NMA_Chemistry_SampleInfo.id)
        ).all()


def _parameters(sample_id: int):
    with session_ctx() as session:
        return session.scalars(
            select(NMA_FieldParameters)
            .where(NMA_FieldParameters.chemistry_sample_info_id == sample_id)
            .order_by(NMA_FieldParameters.field_parameter)
        ).all()


# ------------------------- pure-function tests -------------------------------


def test_prep_sample_info_reads_headings_with_hints():
    prepped = prep_sample_info(_sample_info_row())

    assert prepped["well_pointid"] == WELL
    assert prepped["collection_date"] == datetime(2025, 6, 10, 12, 5)
    assert prepped["attributes"]["sample_type"] == "SD"
    assert prepped["attributes"]["collection_method"] == "F"
    assert prepped["attributes"]["data_source"] == "NMBGMR"


@pytest.mark.parametrize(
    "written,code",
    [
        ("Bailer", "B"),
        ("Faucet at well head", "F"),
        ("Grab sample", "G"),
        ("Faucet or outlet at house", "H"),
        ("Pump", "P"),
        ("Thief sampler", "T"),
        ("Unknown", "U"),
        # Case and spacing are forgiven.
        ("faucet AT  well head ", "F"),
        # The legacy code itself is accepted too.
        ("F", "F"),
        ("h", "H"),
        (" U ", "U"),
    ],
)
def test_collection_method_is_stored_as_its_legacy_code(written, code):
    prepped = prep_sample_info(
        _sample_info_row(**{"CollectionMethod (F = faucet)": written})
    )
    assert prepped["attributes"]["collection_method"] == code


def test_collection_method_outside_the_vocabulary_is_refused():
    with pytest.raises(FieldParamsMappingError) as exc:
        prep_sample_info(_sample_info_row(**{"CollectionMethod (F = faucet)": "X"}))
    assert "not a known collection method" in str(exc.value)


def test_collection_method_may_be_left_blank():
    prepped = prep_sample_info(
        _sample_info_row(**{"CollectionMethod (F = faucet)": None})
    )
    assert "collection_method" not in prepped["attributes"]


def test_prep_sample_info_folds_staff_into_notes():
    """Staff has no legacy column, and the names are worth more than the space."""
    prepped = prep_sample_info(_sample_info_row())

    notes = prepped["attributes"]["sample_notes"]
    assert notes.startswith("Staff: Dan Lavery, Sianin Spaur")
    assert "Sampled from spigot" in notes


def test_prep_sample_info_requires_a_collection_date():
    with pytest.raises(FieldParamsMappingError):
        prep_sample_info(_sample_info_row(CollectionDate=None))


def test_prep_sample_info_rejects_an_unparseable_date():
    with pytest.raises(FieldParamsMappingError):
        prep_sample_info(_sample_info_row(CollectionDate="last Tuesday"))


def test_prep_sample_info_rejects_a_sample_point_from_another_well():
    with pytest.raises(FieldParamsMappingError) as exc:
        prep_sample_info(_sample_info_row(**{"SamplePointID": "Other WellA"}))
    assert "does not belong" in str(exc.value)


def test_prep_sample_info_rejects_an_overlong_collected_by():
    with pytest.raises(FieldParamsMappingError) as exc:
        prep_sample_info(
            _sample_info_row(**{"CollectedBy (GR = grab??)": "Sianin Spaur"})
        )
    assert "5" in str(exc.value)


def test_prep_sample_info_rejects_an_unassigned_pointid():
    """`WL-####` is a real sample whose well has no identifier yet."""
    with pytest.raises(FieldParamsMappingError) as exc:
        prep_sample_info(
            _sample_info_row(WellPointID="WL-####", **{"SamplePointID": "WL-####A"})
        )
    assert "no well identifier assigned" in str(exc.value)


def test_prep_field_parameters_rejects_an_unassigned_pointid():
    with pytest.raises(FieldParamsMappingError):
        prep_field_parameters(_field_params_row(**{"SamplePointID": "WL-####A"}))


@pytest.mark.parametrize(
    "written,expected",
    [
        ("2025-06-10T12:05:00", datetime(2025, 6, 10, 12, 5)),
        # Hand-typed single-digit hour: not ISO, plainly meant as 02:15.
        ("2025-06-06T2:15:00", datetime(2025, 6, 6, 2, 15)),
        ("2026-03-09 14:04:00", datetime(2026, 3, 9, 14, 4)),
        ("6/10/2025", datetime(2025, 6, 10)),
    ],
)
def test_collection_dates_the_sheet_actually_holds(written, expected):
    assert (
        prep_sample_info(_sample_info_row(CollectionDate=written))["collection_date"]
        == expected
    )


def test_prep_field_parameters_unpivots_columns_to_symbols():
    prepped = prep_field_parameters(_field_params_row())

    by_symbol = {m["symbol"]: m for m in prepped["measurements"]}
    assert by_symbol["pHf"]["value"] == 6.94
    assert by_symbol["pHf"]["units"] == "pH"
    assert by_symbol["T"]["units"] == "°C"
    assert by_symbol["CF"]["units"] == "µS/cm"
    assert by_symbol["DO"]["value"] == 1.35
    assert by_symbol["ORP"]["value"] == 78.5
    assert by_symbol["Q"]["units"] == "gpm"


def test_prep_field_parameters_skips_empty_cells():
    prepped = prep_field_parameters(
        _field_params_row(**{"DO (mg/L)": None, "ORP (mV)": ""})
    )
    symbols = {m["symbol"] for m in prepped["measurements"]}
    assert "DO" not in symbols
    assert "ORP" not in symbols


def test_prep_field_parameters_rejects_a_non_numeric_reading():
    with pytest.raises(FieldParamsMappingError) as exc:
        prep_field_parameters(_field_params_row(pHf="meter broken"))
    assert "pHf" in str(exc.value)


def test_select_tables_finds_renamed_tabs_by_their_headers():
    tables = [
        _table("Sheet1", SAMPLE_INFO_HEADER, [_sample_info_row()]),
        _table("Sheet2", FIELD_PARAMS_HEADER, [_field_params_row()]),
    ]
    sample_info, field_params = select_tables(tables)

    assert sample_info.title == "Sheet1"
    assert field_params.title == "Sheet2"


def test_select_tables_ignores_lab_result_tabs():
    tables = _tables() + [
        _table("GenChemResults", ["WellPointID", "AnalysisDate", "Ca (mg/L)"], [])
    ]
    sample_info, field_params = select_tables(tables)

    assert sample_info.title == "ChemistrySampleInfo"
    assert field_params.title == "FieldParameters"


# ------------------------- ingestion tests -----------------------------------


def test_import_creates_sample_and_field_parameters(
    water_well_thing, _cleanup_field_chemistry
):
    result = import_field_tables(_tables())

    assert result.exit_code == 0, result.stderr
    (sample,) = _samples()
    assert sample.nma_sample_point_id == f"{WELL}A"
    assert sample.collection_date == datetime(2025, 6, 10, 12, 5)
    assert sample.sample_type == "SD"
    assert sample.collection_method == "F"

    parameters = _parameters(sample.id)
    assert {p.field_parameter for p in parameters} == {
        "pHf",
        "T",
        "CF",
        "DO",
        "ORP",
        "Q",
    }
    assert result.payload["summary"]["total_rows_imported"] == 6
    assert result.payload["summary"]["samples_created"] == 1


def test_measurement_time_is_kept_in_the_notes(
    water_well_thing, _cleanup_field_chemistry
):
    import_field_tables(_tables())

    (sample,) = _samples()
    assert all("2025-06-10T12:07:00" in p.notes for p in _parameters(sample.id))


def test_rerunning_loads_nothing_twice(water_well_thing, _cleanup_field_chemistry):
    import_field_tables(_tables())
    second = import_field_tables(_tables())

    assert second.exit_code == 0, second.stderr
    assert len(_samples()) == 1
    (sample,) = _samples()
    assert len(_parameters(sample.id)) == 6
    assert second.payload["summary"]["total_rows_imported"] == 0
    assert second.payload["summary"]["samples_matched"] == 1
    assert second.payload["summary"]["parameters_skipped"] == 6


def test_a_second_visit_becomes_its_own_sample(
    water_well_thing, _cleanup_field_chemistry
):
    import_field_tables(_tables())
    import_field_tables(
        _tables(
            sample_rows=[
                _sample_info_row(
                    CollectionDate="2025-09-02T09:00:00",
                    **{"SamplePointID": f"{WELL}B"},
                )
            ],
            param_rows=[_field_params_row(**{"SamplePointID": f"{WELL}B"})],
        )
    )

    points = [s.nma_sample_point_id for s in _samples()]
    assert points == [f"{WELL}A", f"{WELL}B"]


def test_a_sample_the_lab_ingest_already_created_is_reused(
    water_well_thing, _cleanup_field_chemistry
):
    """The lab batch and the field visit are one sample, matched on well + date."""
    with session_ctx() as session:
        session.add(
            NMA_Chemistry_SampleInfo(
                thing_id=water_well_thing.id,
                nma_sample_point_id=f"{WELL}A",
                nma_wclab_id="LAB-1",
                collection_date=datetime(2025, 6, 10, 12, 5),
            )
        )
        session.commit()

    result = import_field_tables(_tables())

    assert result.exit_code == 0, result.stderr
    (sample,) = _samples()
    assert sample.nma_wclab_id == "LAB-1"
    assert result.payload["summary"]["samples_created"] == 0
    assert result.payload["summary"]["samples_matched"] == 1
    # The field readings hang off the lab's sample, and carry its lab id.
    assert {p.nma_wclab_id for p in _parameters(sample.id)} == {"LAB-1"}


def test_matching_is_by_calendar_day_when_the_times_differ(
    water_well_thing, _cleanup_field_chemistry
):
    """The sheet's time and the lab's time for one visit differ by minutes."""
    with session_ctx() as session:
        session.add(
            NMA_Chemistry_SampleInfo(
                thing_id=water_well_thing.id,
                nma_sample_point_id=f"{WELL}A",
                nma_wclab_id="LAB-1",
                collection_date=datetime(2025, 6, 10, 14, 30),
            )
        )
        session.commit()

    result = import_field_tables(_tables())

    assert len(_samples()) == 1
    assert result.payload["summary"]["samples_matched"] == 1


def test_existing_values_are_not_overwritten_but_are_reported(
    water_well_thing, _cleanup_field_chemistry
):
    with session_ctx() as session:
        session.add(
            NMA_Chemistry_SampleInfo(
                thing_id=water_well_thing.id,
                nma_sample_point_id=f"{WELL}A",
                collection_date=datetime(2025, 6, 10, 12, 5),
                sample_type="GW",
            )
        )
        session.commit()

    result = import_field_tables(_tables())

    (sample,) = _samples()
    assert sample.sample_type == "GW"
    assert any("sample_type" in w for w in result.payload["warnings"])


def test_a_legacy_collection_method_code_agrees_with_its_meaning(
    water_well_thing, _cleanup_field_chemistry
):
    """Legacy rows hold "F"; the sheet says "Faucet at well head". Same thing."""
    with session_ctx() as session:
        session.add(
            NMA_Chemistry_SampleInfo(
                thing_id=water_well_thing.id,
                nma_sample_point_id=f"{WELL}A",
                collection_date=datetime(2025, 6, 10, 12, 5),
                collection_method="F",
            )
        )
        session.commit()

    result = import_field_tables(_tables())

    assert result.exit_code == 0, result.stderr
    (sample,) = _samples()
    assert sample.collection_method == "F"
    assert not any("collection_method" in w for w in result.payload["warnings"])


def test_unknown_well_aborts_the_whole_import(
    water_well_thing, _cleanup_field_chemistry
):
    result = import_field_tables(
        _tables(
            sample_rows=[
                _sample_info_row(),
                _sample_info_row(
                    WellPointID="No Such Well", **{"SamplePointID": "No Such WellA"}
                ),
            ]
        )
    )

    assert result.exit_code == 1
    assert any("No Such Well" in e for e in result.payload["validation_errors"])
    # Nothing is written, not even the row that was fine.
    assert _samples() == []


def test_field_parameters_for_an_unknown_sample_abort_the_import(
    water_well_thing, _cleanup_field_chemistry
):
    result = import_field_tables(
        _tables(param_rows=[_field_params_row(**{"SamplePointID": "Test WellZ"})])
    )

    assert result.exit_code == 1
    assert any("Test WellZ" in e for e in result.payload["validation_errors"])
    assert _samples() == []


def test_dry_run_writes_nothing(water_well_thing, _cleanup_field_chemistry):
    result = import_field_tables(_tables(), dry_run=True)

    assert result.exit_code == 0, result.stderr
    assert result.payload["summary"]["dry_run"] is True
    assert result.payload["summary"]["samples_created"] == 1
    assert _samples() == []


def test_upload_from_a_downloaded_workbook(
    tmp_path, water_well_thing, _cleanup_field_chemistry
):
    workbook = Workbook()
    workbook.remove(workbook.active)
    info = workbook.create_sheet("ChemistrySampleInfo")
    info.append(SAMPLE_INFO_HEADER)
    info.append([_sample_info_row().get(c) for c in SAMPLE_INFO_HEADER])
    params = workbook.create_sheet("FieldParameters")
    params.append(FIELD_PARAMS_HEADER)
    params.append([_field_params_row().get(c) for c in FIELD_PARAMS_HEADER])
    path = tmp_path / "field.xlsx"
    workbook.save(path)

    # Archiving is covered below; this is about reading the file.
    result = upload_field_export([path], archive=False)

    assert result.exit_code == 0, result.stderr
    (sample,) = _samples()
    assert len(_parameters(sample.id)) == 6


def test_upload_from_one_csv_per_tab(
    tmp_path, water_well_thing, _cleanup_field_chemistry
):
    info_path = tmp_path / "info.csv"
    info_path.write_text(
        ",".join(SAMPLE_INFO_HEADER)
        + "\n"
        + ",".join(
            str(_sample_info_row().get(c) or "") for c in SAMPLE_INFO_HEADER
        ).replace("Dan Lavery, Sianin Spaur", "Dan Lavery")
        + "\n"
    )
    params_path = tmp_path / "params.csv"
    params_path.write_text(
        ",".join(FIELD_PARAMS_HEADER)
        + "\n"
        + ",".join(str(_field_params_row().get(c) or "") for c in FIELD_PARAMS_HEADER)
        + "\n"
    )

    # Each CSV holds one tab, so both are read and pooled.
    tables = read_local_export(info_path) + read_local_export(params_path)
    sample_info, field_params = select_tables(tables)
    assert sample_info is not None and field_params is not None

    result = upload_field_export([info_path, params_path], archive=False)

    assert result.exit_code == 0, result.stderr
    (sample,) = _samples()
    assert len(_parameters(sample.id)) == 6


# ------------------------- raw zone -------------------------------------------


@pytest.fixture()
def raw_zone(tmp_path, monkeypatch):
    """A raw zone on disk, with dlt's own state kept out of the developer's home."""
    monkeypatch.setenv("DLT_DATA_DIR", str(tmp_path / "dlt"))
    return (tmp_path / "raw").resolve().as_uri()


def _workbook(tmp_path, sample_rows=None, param_rows=None):
    workbook = Workbook()
    workbook.remove(workbook.active)
    info = workbook.create_sheet("ChemistrySampleInfo")
    info.append(SAMPLE_INFO_HEADER)
    for row in sample_rows if sample_rows is not None else [_sample_info_row()]:
        info.append([row.get(c) for c in SAMPLE_INFO_HEADER])
    params = workbook.create_sheet("FieldParameters")
    params.append(FIELD_PARAMS_HEADER)
    for row in param_rows if param_rows is not None else [_field_params_row()]:
        params.append([row.get(c) for c in FIELD_PARAMS_HEADER])
    path = tmp_path / "field.xlsx"
    workbook.save(path)
    return path


def test_upload_archives_what_it_read_then_loads_that(
    tmp_path, raw_zone, water_well_thing, _cleanup_field_chemistry
):
    """What gets loaded is what was archived, not a second read of the source."""
    result = upload_field_export([_workbook(tmp_path)], raw_url=raw_zone)

    assert result.exit_code == 0, result.stderr
    raw = result.payload["raw"]
    assert raw["load_id"]
    assert raw["dataset"] == FIELD_SHEET_DATASET
    assert raw["rows_archived"] == 2

    (sample,) = _samples()
    assert len(_parameters(sample.id)) == 6


def test_replay_reloads_a_snapshot_without_the_source(
    tmp_path, raw_zone, water_well_thing, _cleanup_field_chemistry
):
    """A mapping fix is retested against the rows the original run saw."""
    path = _workbook(tmp_path)
    first = upload_field_export([path], raw_url=raw_zone)
    load_id = first.payload["raw"]["load_id"]

    # The source is gone; the archive is not.
    path.unlink()
    with session_ctx() as session:
        session.execute(
            delete(NMA_Chemistry_SampleInfo).where(
                NMA_Chemistry_SampleInfo.nma_sample_point_id.like(f"{WELL}%")
            )
        )
        session.commit()

    replayed = replay_field_sheet(load_id, raw_url=raw_zone)

    assert replayed.exit_code == 0, replayed.stderr
    (sample,) = _samples()
    assert len(_parameters(sample.id)) == 6
    assert replayed.payload["raw"]["replayed"] is True


def test_replay_is_still_idempotent(
    tmp_path, raw_zone, water_well_thing, _cleanup_field_chemistry
):
    load_id = upload_field_export([_workbook(tmp_path)], raw_url=raw_zone).payload[
        "raw"
    ]["load_id"]

    again = replay_field_sheet(load_id, raw_url=raw_zone)

    assert again.payload["summary"]["total_rows_imported"] == 0
    assert again.payload["summary"]["parameters_skipped"] == 6
    assert len(_samples()) == 1


def test_a_dry_run_leaves_no_snapshot_behind(
    tmp_path, raw_zone, water_well_thing, _cleanup_field_chemistry
):
    """Checking what would happen should not write to the archive either."""
    result = upload_field_export([_workbook(tmp_path)], dry_run=True, raw_url=raw_zone)

    assert result.exit_code == 0, result.stderr
    assert result.payload["raw"] == {}
    assert list_snapshots(FIELD_SHEET_DATASET, raw_url=raw_zone) == []
    assert _samples() == []


def test_archiving_can_be_skipped(
    tmp_path, raw_zone, water_well_thing, _cleanup_field_chemistry
):
    result = upload_field_export([_workbook(tmp_path)], archive=False, raw_url=raw_zone)

    assert result.exit_code == 0, result.stderr
    assert result.payload["raw"] == {}
    assert list_snapshots(FIELD_SHEET_DATASET, raw_url=raw_zone) == []
    assert len(_samples()) == 1


def test_the_archive_keeps_a_bad_row_as_written(
    tmp_path, raw_zone, water_well_thing, _cleanup_field_chemistry
):
    """A run that aborts still archives: that is the record of what arrived."""
    path = _workbook(
        tmp_path, sample_rows=[_sample_info_row(CollectionDate=None)], param_rows=[]
    )
    result = upload_field_export([path], raw_url=raw_zone)

    assert result.exit_code == 1
    assert result.payload["raw"]["load_id"]
    (archived,) = read_snapshot(FIELD_SHEET_DATASET, raw_url=raw_zone)
    assert archived.rows[0]["CollectionDate"] is None


# ============= EOF =============================================
