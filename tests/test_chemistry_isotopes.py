# ===============================================================================
# Copyright 2026 ross
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
# ===============================================================================
"""Tests for the isotope results ingest (services/chemistry_isotopes.py)."""

import re
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, select

from db.engine import session_ctx
from db.nma_legacy import NMA_Chemistry_SampleInfo, NMA_MinorTraceChemistry
from services.chemistry_field_sheet import SheetTable
from services.chemistry_isotopes import (
    ISOTOPE_COLUMNS,
    ISOTOPE_DATASET,
    IsotopeMappingError,
    import_isotope_table,
    prep_isotope_row,
    replay_isotope_sheet,
    select_isotope_table,
    upload_isotope_export,
)
from services.chemistry_lims import bulk_upload_chemistry
from services.legacy_chemistry import LEGACY_UNITS

WELL = "Test Well"
POINT = f"{WELL}A"
COLLECTED = datetime(2025, 6, 1, 10, 0)

HEADER = [
    "WellPointID",
    "AnalysisDate",
    "AnalysisLab",
    "ChemistryLabID",
    "Age_Qualifier",
    "Age (years)",
    "Age_Uncertainty (years)",
    "C13r (pm)",
    "H2r (pm)",
    "H2r_Uncertainty (pm)",
    "O18r (pm)",
    "O18r_Uncertainty (pm)",
    "pMC_Qualifier",
    "pMC (pmc)",
    "pMC_Uncertainty (pmc)",
    "Tritium Units (TU)",
    "Tritium_error",
]


def _beta_row(**overrides):
    row = {
        "WellPointID": POINT,
        "AnalysisDate": datetime(2025, 8, 27),
        "AnalysisLab": "Beta Analytic",
        # Typed as a number in the sheet, so openpyxl reads it back as a float.
        "ChemistryLabID": 765646.0,
        "Age (years)": 10885.0,
        "Age_Uncertainty (years)": 30.0,
        "C13r (pm)": -8.79,
        "H2r (pm)": -100.58,
        "O18r (pm)": -13.55,
        "pMC (pmc)": 25.79,
        "pMC_Uncertainty (pmc)": 0.09,
    }
    row.update(overrides)
    return row


def _nmbgmr_row(**overrides):
    row = {
        "WellPointID": POINT,
        "AnalysisLab": "NMBGMR",
        "H2r (pm)": -103.365,
        "H2r_Uncertainty (pm)": 0.3,
        "O18r (pm)": -13.963,
        "O18r_Uncertainty (pm)": 0.05,
    }
    row.update(overrides)
    return row


def _miami_row(**overrides):
    row = {
        "WellPointID": POINT,
        "AnalysisDate": datetime(2025, 6, 12),
        "AnalysisLab": "University of Miami Tritium Laboratory",
        "Tritium Units (TU)": 1.39,
        "Tritium_error": 0.09,
    }
    row.update(overrides)
    return row


def _record(row: dict, row_number: int = 2) -> dict:
    record = {column: row.get(column) for column in HEADER}
    record[SheetTable.ROW_NUMBER_KEY] = row_number
    return record


def _table(rows: list[dict], title: str = "IsotopeResults") -> SheetTable:
    return SheetTable(
        title=title,
        header=list(HEADER),
        rows=[_record(row, offset + 2) for offset, row in enumerate(rows)],
    )


def _workbook(tmp_path, rows: list[dict]):
    """The spreadsheet as downloaded: the isotope tab beside the others."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.create_sheet("ChemistrySampleInfo").append(["WellPointID"])
    sheet = workbook.create_sheet("IsotopeResults")
    sheet.append(HEADER)
    for row in rows:
        sheet.append([row.get(c) for c in HEADER])
    path = tmp_path / "field.xlsx"
    workbook.save(path)
    return path


def _record_sample(thing_id: int, point: str = POINT, when=COLLECTED, wclab_id=None):
    with session_ctx() as session:
        sample = NMA_Chemistry_SampleInfo(
            thing_id=thing_id,
            nma_sample_point_id=point,
            nma_wclab_id=wclab_id,
            collection_date=when,
        )
        session.add(sample)
        session.commit()
        return sample.id


def _sample(sample_id: int) -> NMA_Chemistry_SampleInfo:
    with session_ctx() as session:
        return session.get(NMA_Chemistry_SampleInfo, sample_id)


def _results(sample_id: int):
    with session_ctx() as session:
        return session.scalars(
            select(NMA_MinorTraceChemistry)
            .where(NMA_MinorTraceChemistry.chemistry_sample_info_id == sample_id)
            .order_by(
                NMA_MinorTraceChemistry.analyses_agency,
                NMA_MinorTraceChemistry.analyte,
            )
        ).all()


@pytest.fixture()
def _cleanup_isotopes():
    """Remove samples created during a test; their results cascade with them."""
    yield
    with session_ctx() as session:
        session.execute(
            delete(NMA_Chemistry_SampleInfo).where(
                NMA_Chemistry_SampleInfo.nma_sample_point_id.like(f"{WELL}%")
            )
        )
        session.commit()


# ------------------------- pure-function tests -------------------------------


def test_a_beta_row_unpivots_to_one_result_per_measurement():
    row = prep_isotope_row(
        _record(
            _beta_row(
                **{
                    "Age_Qualifier": ">",
                    "Age (years)": 43500.0,
                    "Age_Uncertainty (years)": None,
                    "pMC_Qualifier": "<",
                    "pMC (pmc)": 0.44,
                    "pMC_Uncertainty (pmc)": None,
                }
            )
        )
    )

    assert (row["sample_point_id"], row["well_pointid"], row["lab"]) == (
        POINT,
        WELL,
        "Beta Analytic",
    )
    assert row["analysis_date"] == datetime(2025, 8, 27)
    assert {
        r["analyte"]: (r["symbol"], r["value"], r["units"]) for r in row["results"]
    } == {
        "C14_years": (">", 43500.0, "yrs"),
        "C14": ("<", 0.44, "pmc"),
        "C13r": (None, -8.79, "pm"),
        "H2r": (None, -100.58, "pm"),
        "O18r": (None, -13.55, "pm"),
    }


def test_uncertainties_go_with_their_measurement():
    row = prep_isotope_row(_record(_nmbgmr_row()))

    assert {r["analyte"]: r["uncertainty"] for r in row["results"]} == {
        "H2r": 0.3,
        "O18r": 0.05,
    }


def test_blank_qualifiers_are_normal():
    row = prep_isotope_row(_record(_beta_row()))

    assert all(r["symbol"] is None for r in row["results"])


@pytest.mark.parametrize(
    ("cell", "expected"),
    [
        (765646.0, "765646"),
        (765646, "765646"),
        ("  885-44878-1 ", "885-44878-1"),
        ("260404-01\n", "260404-01"),
        (None, None),
    ],
)
def test_a_lab_id_is_stored_as_text(cell, expected):
    row = prep_isotope_row(_record(_beta_row(ChemistryLabID=cell)))

    assert row["lab_id"] == expected


def test_a_lab_is_matched_regardless_of_case_and_spacing():
    row = prep_isotope_row(_record(_beta_row(AnalysisLab="  beta   ANALYTIC ")))

    assert row["lab"] == "Beta Analytic"


@pytest.mark.parametrize(
    ("overrides", "message"),
    [
        ({"WellPointID": None}, "Missing SamplePointID"),
        ({"WellPointID": WELL}, "has no sample letter"),
        ({"WellPointID": "WL-####A"}, "no well identifier assigned"),
        ({"AnalysisLab": None}, "Missing AnalysisLab"),
        ({"AnalysisLab": "Beta Analytics"}, "is not a known lab"),
        ({"AnalysisDate": "not a date"}, "is not a recognizable date"),
        ({"H2r (pm)": "n/a"}, "H2r (pm)='n/a' is not a number"),
        # The qualifier has its own column; one typed into the value is a slip.
        ({"pMC (pmc)": "<0.44"}, "pMC (pmc)='<0.44' is not a number"),
        ({"Age_Qualifier": "~"}, "Age_Qualifier='~' is not one of"),
        (
            {
                "Age (years)": None,
                "Age_Qualifier": ">",
                "Age_Uncertainty (years)": None,
            },
            "Age_Qualifier filled but Age (years) blank",
        ),
        (
            {"pMC (pmc)": None},
            "pMC_Uncertainty (pmc) filled but pMC (pmc) blank",
        ),
        ({"ChemistryLabID": "X" * 26}, "longer than 25 characters"),
    ],
)
def test_a_row_that_cannot_load_is_refused(overrides, message):
    with pytest.raises(IsotopeMappingError, match=re.escape(message)):
        prep_isotope_row(_record(_beta_row(**overrides)))


def test_every_problem_in_a_row_is_named_at_once():
    with pytest.raises(IsotopeMappingError) as excinfo:
        prep_isotope_row(_record(_beta_row(**{"H2r (pm)": "x", "O18r (pm)": "y"})))

    assert "H2r (pm)='x'" in str(excinfo.value)
    assert "O18r (pm)='y'" in str(excinfo.value)


def test_a_row_without_results_is_empty_not_wrong():
    row = prep_isotope_row(
        _record({"WellPointID": POINT, "AnalysisLab": "Beta Analytic"})
    )

    assert row["results"] == []


def test_headings_match_with_or_without_their_hints():
    record = {
        "SamplePointID": POINT,
        "AnalysisLab": "NMBGMR",
        "H2r": -103.365,
        "H2r_Uncertainty": 0.3,
        "Tritium Units": 1.39,
        SheetTable.ROW_NUMBER_KEY: 2,
    }

    row = prep_isotope_row(record)

    assert {r["analyte"]: r["value"] for r in row["results"]} == {
        "H2r": -103.365,
        "3H": 1.39,
    }


def test_units_are_the_legacy_abbreviations():
    assert {column.units for column in ISOTOPE_COLUMNS} <= set(LEGACY_UNITS)


def test_the_isotope_tab_is_found_by_title_or_by_its_columns():
    general = SheetTable(
        title="GenChemResults",
        header=["WellPointID", "AnalysisLab", "Ca (mg/L)"],
    )
    renamed = _table([_beta_row()], title="isotopes (copy)")

    assert select_isotope_table([general, _table([])]).title == "IsotopeResults"
    assert select_isotope_table([general, renamed]) is renamed
    assert select_isotope_table([general]) is None


# ------------------------- import tests --------------------------------------


def test_results_from_three_labs_attach_to_their_sample(
    water_well_thing, _cleanup_isotopes
):
    sample_id = _record_sample(water_well_thing.id)

    result = import_isotope_table([_table([_beta_row(), _miami_row(), _nmbgmr_row()])])

    assert result.exit_code == 0, result.stderr
    assert result.payload["summary"]["total_rows_imported"] == 8
    stored = _results(sample_id)
    assert [(r.analyses_agency, r.analyte) for r in stored] == [
        ("Beta Analytic", "C13r"),
        ("Beta Analytic", "C14"),
        ("Beta Analytic", "C14_years"),
        ("Beta Analytic", "H2r"),
        ("Beta Analytic", "O18r"),
        ("NMBGMR", "H2r"),
        ("NMBGMR", "O18r"),
        ("University of Miami Tritium Laboratory", "3H"),
    ]
    by_key = {(r.analyses_agency, r.analyte): r for r in stored}
    # Two labs' deuterium, side by side.
    assert by_key[("Beta Analytic", "H2r")].sample_value == -100.58
    assert by_key[("NMBGMR", "H2r")].sample_value == -103.365
    assert by_key[("NMBGMR", "H2r")].uncertainty == 0.3
    # Each lab's id on its own results.
    assert {r.analyses_agency: r.nma_wclab_id for r in stored} == {
        "Beta Analytic": "765646",
        "NMBGMR": None,
        "University of Miami Tritium Laboratory": None,
    }
    assert by_key[("Beta Analytic", "C14")].analysis_date == date(2025, 8, 27)
    for row in stored:
        assert row.nma_sample_point_id == POINT
        # Legacy audit columns hold NM_Aquifer's values only.
        assert row.nma_global_id is None
        assert row.nma_chemistry_sample_info_uuid is None


def test_the_sample_record_is_left_alone(water_well_thing, _cleanup_isotopes):
    """Its one lab-id slot is the LIMS ingest's, which adopts only an empty one."""
    sample_id = _record_sample(water_well_thing.id)

    import_isotope_table([_table([_beta_row()])])

    sample = _sample(sample_id)
    assert sample.nma_wclab_id is None
    assert sample.analyses_agency is None


def test_rerunning_loads_nothing_twice(water_well_thing, _cleanup_isotopes):
    sample_id = _record_sample(water_well_thing.id)
    tables = [_table([_beta_row(), _nmbgmr_row()])]

    import_isotope_table(tables)
    again = import_isotope_table(tables)

    assert again.exit_code == 0, again.stderr
    assert again.payload["summary"]["total_rows_imported"] == 0
    assert again.payload["summary"]["results_skipped"] == 7
    assert len(_results(sample_id)) == 7


def test_a_changed_result_stops_the_run_for_review(water_well_thing, _cleanup_isotopes):
    sample_id = _record_sample(water_well_thing.id)
    import_isotope_table([_table([_beta_row()])])

    # A re-analysis, or a typo: either way a person decides which stands.
    changed = import_isotope_table(
        [_table([_beta_row(**{"H2r (pm)": -99.0}), _nmbgmr_row()])]
    )

    assert changed.exit_code == 1
    (failure,) = changed.payload["failed_rows"]
    assert failure["row"] == 2
    assert "H2r from Beta Analytic is already stored" in failure["reason"]
    assert "value -100.58 stored, -99.0 here" in failure["reason"]
    stored = _results(sample_id)
    # Nothing written: not the change, and not the other row's new results.
    assert len(stored) == 5
    assert {r.analyte: r.sample_value for r in stored}["H2r"] == -100.58


def test_a_changed_qualifier_counts_as_a_change(water_well_thing, _cleanup_isotopes):
    _record_sample(water_well_thing.id)
    import_isotope_table([_table([_beta_row()])])

    changed = import_isotope_table([_table([_beta_row(pMC_Qualifier="<")])])

    assert changed.exit_code == 1
    assert (
        "qualifier None stored, '<' here" in changed.payload["failed_rows"][0]["reason"]
    )


def test_a_sample_not_yet_recorded_fails_its_row(water_well_thing, _cleanup_isotopes):
    result = import_isotope_table([_table([_beta_row()])])

    assert result.exit_code == 1
    (failure,) = result.payload["failed_rows"]
    assert failure["sample_point_id"] == POINT
    assert f"no sample {POINT}" in failure["reason"]
    assert "field-upload" in failure["reason"]


def test_an_unknown_well_fails_its_row(water_well_thing, _cleanup_isotopes):
    result = import_isotope_table([_table([_beta_row(WellPointID="ZZ-9999A")])])

    assert result.exit_code == 1
    assert (
        "no matching Thing (well) ZZ-9999" in result.payload["failed_rows"][0]["reason"]
    )


def test_two_samples_sharing_the_id_fail_the_row(water_well_thing, _cleanup_isotopes):
    _record_sample(water_well_thing.id)
    _record_sample(water_well_thing.id, when=datetime(2025, 7, 1))

    result = import_isotope_table([_table([_beta_row()])])

    assert result.exit_code == 1
    assert "cannot tell which one" in result.payload["failed_rows"][0]["reason"]


def test_results_cannot_predate_collection(water_well_thing, _cleanup_isotopes):
    sample_id = _record_sample(water_well_thing.id)

    result = import_isotope_table(
        [_table([_miami_row(AnalysisDate=datetime(2025, 5, 30))])]
    )

    assert result.exit_code == 1
    reason = result.payload["failed_rows"][0]["reason"]
    assert "collected on 2025-06-01" in reason
    assert "analysed it on 2025-05-30" in reason
    assert _results(sample_id) == []


def test_an_undated_sample_loads_with_a_warning(water_well_thing, _cleanup_isotopes):
    sample_id = _record_sample(water_well_thing.id, when=None)

    result = import_isotope_table([_table([_nmbgmr_row()])])

    assert result.exit_code == 0, result.stderr
    assert "has no collection date" in result.payload["warnings"][0]
    assert len(_results(sample_id)) == 2


def test_two_different_results_in_the_sheet_fail_both_rows(
    water_well_thing, _cleanup_isotopes
):
    sample_id = _record_sample(water_well_thing.id)

    result = import_isotope_table(
        [_table([_nmbgmr_row(), _nmbgmr_row(**{"H2r (pm)": -50.0})])]
    )

    assert result.exit_code == 1
    assert sorted(f["row"] for f in result.payload["failed_rows"]) == [2, 3]
    assert _results(sample_id) == []


def test_a_repeated_row_loads_once(water_well_thing, _cleanup_isotopes):
    sample_id = _record_sample(water_well_thing.id)

    result = import_isotope_table([_table([_nmbgmr_row(), _nmbgmr_row()])])

    assert result.exit_code == 0, result.stderr
    assert "repeats row 2; loaded once" in result.payload["warnings"][0]
    assert len(_results(sample_id)) == 2


def test_a_column_the_ingest_does_not_read_is_named(
    water_well_thing, _cleanup_isotopes
):
    """A lab's new measurement would otherwise never load, with no sign why."""
    sample_id = _record_sample(water_well_thing.id)
    table = _table([_nmbgmr_row()])
    table.header += ["87Sr/86Sr", ""]
    for record in table.rows:
        record["87Sr/86Sr"] = 0.7101

    result = import_isotope_table([table])

    assert result.exit_code == 0, result.stderr
    (warning,) = result.payload["warnings"]
    assert "doesn't read the column(s) '87Sr/86Sr'" in warning
    assert len(_results(sample_id)) == 2


def test_the_template_columns_raise_no_warning(water_well_thing, _cleanup_isotopes):
    _record_sample(water_well_thing.id)

    result = import_isotope_table([_table([_beta_row()])])

    assert result.payload["warnings"] == []


def test_empty_rows_are_reported_and_skipped(water_well_thing, _cleanup_isotopes):
    _record_sample(water_well_thing.id)

    result = import_isotope_table(
        [_table([{"WellPointID": POINT, "AnalysisLab": "NMBGMR"}, _miami_row()])]
    )

    assert result.exit_code == 0, result.stderr
    assert result.payload["empty_rows"] == [
        {"row": 2, "sample_point_id": POINT, "analysis_lab": "NMBGMR"}
    ]
    assert result.payload["summary"]["total_rows_imported"] == 1


def test_one_bad_row_writes_nothing(water_well_thing, _cleanup_isotopes):
    sample_id = _record_sample(water_well_thing.id)

    result = import_isotope_table(
        [_table([_beta_row(), _nmbgmr_row(AnalysisLab="Somewhere Else")])]
    )

    assert result.exit_code == 1
    assert result.payload["summary"]["total_rows_imported"] == 0
    assert _results(sample_id) == []


def test_a_dry_run_writes_nothing(water_well_thing, _cleanup_isotopes):
    sample_id = _record_sample(water_well_thing.id)

    result = import_isotope_table([_table([_beta_row()])], dry_run=True)

    assert result.exit_code == 0, result.stderr
    assert result.payload["summary"]["dry_run"] is True
    assert result.payload["summary"]["total_rows_imported"] == 5
    assert _results(sample_id) == []


def test_a_workbook_without_the_tab_is_refused():
    result = import_isotope_table([SheetTable(title="ChemistrySampleInfo")])

    assert result.exit_code == 1
    assert "Found no IsotopeResults tab" in result.payload["validation_errors"][0]


# ------------------------- reports and the raw zone --------------------------


def test_failed_rows_and_a_loadable_copy_are_written(
    tmp_path, water_well_thing, _cleanup_isotopes
):
    sample_id = _record_sample(water_well_thing.id)
    path = _workbook(tmp_path, [_beta_row(), _nmbgmr_row(WellPointID=WELL)])

    result = upload_isotope_export(
        [path],
        dry_run=True,
        failed_rows_path=tmp_path / "failed.xlsx",
        loadable_path=tmp_path / "loadable.xlsx",
    )

    assert result.exit_code == 1
    failed = load_workbook(result.payload["reports"]["failed_rows"])
    sheet = failed["IsotopeResults"]
    assert sheet["A2"].value == 3
    assert "has no sample letter" in sheet["B2"].value

    # The loadable copy holds the rest, and loads.
    loadable = upload_isotope_export(
        [result.payload["reports"]["loadable"]], archive=False
    )
    assert loadable.exit_code == 0, loadable.stderr
    assert len(_results(sample_id)) == 5


def test_upload_archives_the_tab_then_loads_that(
    tmp_path, raw_zone, water_well_thing, _cleanup_isotopes
):
    sample_id = _record_sample(water_well_thing.id)

    result = upload_isotope_export(
        [_workbook(tmp_path, [_beta_row(), _miami_row()])], raw_url=raw_zone
    )

    assert result.exit_code == 0, result.stderr
    raw = result.payload["raw"]
    assert raw["dataset"] == ISOTOPE_DATASET
    # Only the isotope tab: the field sheet archives its own tabs.
    assert raw["rows_archived"] == 2
    assert len(_results(sample_id)) == 6


def test_replay_reloads_a_snapshot_and_stays_idempotent(
    tmp_path, raw_zone, water_well_thing, _cleanup_isotopes
):
    sample_id = _record_sample(water_well_thing.id)
    path = _workbook(tmp_path, [_beta_row()])
    load_id = upload_isotope_export([path], raw_url=raw_zone).payload["raw"]["load_id"]
    path.unlink()

    again = replay_isotope_sheet(load_id, raw_url=raw_zone)

    assert again.exit_code == 0, again.stderr
    assert again.payload["raw"]["replayed"] is True
    assert again.payload["summary"]["results_skipped"] == 5
    assert len(_results(sample_id)) == 5


# ------------------------- with the LIMS ingest ------------------------------


def test_lims_still_adopts_the_sample_after_isotopes_load(
    tmp_path, water_well_thing, _cleanup_isotopes
):
    """Isotope results leave the sample's lab-id slot for Bonnie's lab."""
    sample_id = _record_sample(water_well_thing.id)
    assert import_isotope_table([_table([_beta_row()])]).exit_code == 0

    workbook = Workbook()
    header = [
        "Param",
        "Results_Units",
        "Dilution",
        "AnalysisTime",
        "SampleNumber",
        "SamplePointID",
        "Method",
        "ReportedND",
        "LowerLimit",
        "SampleDate",
    ]
    workbook.active.append(header)
    workbook.active.append(
        [
            "lithium",
            "mg/L",
            1,
            "2025-06-15",
            "LAB-1",
            POINT,
            "EPA 200.7",
            "0.05",
            0.01,
            "2025-06-01",
        ]
    )
    path = tmp_path / "lims.xlsx"
    workbook.save(path)

    lims = bulk_upload_chemistry(path)

    assert lims.exit_code == 0, lims.stderr
    assert lims.payload["summary"]["samples_adopted"] == 1
    assert _sample(sample_id).nma_wclab_id == "LAB-1"
    analytes = [r.analyte for r in _results(sample_id)]
    assert analytes.count("Li") == 1
    assert {"C13r", "C14", "C14_years", "H2r", "O18r"} <= set(analytes)
